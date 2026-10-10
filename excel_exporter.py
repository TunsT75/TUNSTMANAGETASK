import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from config import EXCEL_PATH
from database import get_all_tasks, get_all_work_logs, get_daily_checklist, get_now

def export_to_excel():
    """Tạo hoặc cập nhật file Excel Quản Lý Công Việc bắt đầu ngay từ Dòng 1"""
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # STYLES DEFINITION
    # -------------------------------------------------------------
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark Navy
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    data_font = Font(name="Segoe UI", size=10)
    bold_data_font = Font(name="Segoe UI", size=10, bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    
    status_fills = {
        "Hoàn thành": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"), # Soft Green
        "Chờ làm": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),    # Soft Yellow
        "Đang làm": PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"),   # Soft Blue
        "Đang nhắc": PatternFill(start_color="FDE047", end_color="FDE047", fill_type="solid"),  # Bright Yellow
        "Đã hủy": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")     # Soft Gray
    }
    status_fonts = {
        "Hoàn thành": Font(name="Segoe UI", size=10, bold=True, color="166534"),
        "Chờ làm": Font(name="Segoe UI", size=10, bold=True, color="92400E"),
        "Đang làm": Font(name="Segoe UI", size=10, bold=True, color="1E40AF"),
        "Đang nhắc": Font(name="Segoe UI", size=10, bold=True, color="854D0E"),
        "Đã hủy": Font(name="Segoe UI", size=10, color="475569")
    }
    
    priority_fonts = {
        "Cao": Font(name="Segoe UI", size=10, bold=True, color="DC2626"),
        "Trung bình": Font(name="Segoe UI", size=10, color="D97706"),
        "Thấp": Font(name="Segoe UI", size=10, color="2563EB")
    }

    # -------------------------------------------------------------
    # TAB 1: CHECKLIST HÀNG NGÀY (TODAY'S DAILY CHECKLIST)
    # -------------------------------------------------------------
    ws0 = wb.active
    ws0.title = "📆 Checklist hôm nay"
    ws0.views.sheetView[0].showGridLines = True
    
    headers0 = ["STT", "Checkbox", "Tên công việc", "Loại việc", "Phân loại", "Độ ưu tiên", "Hạn chót / Giờ nhắc", "Trạng thái", "Giờ hoàn thành"]
    for col_num, header in enumerate(headers0, 1):
        cell = ws0.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    daily_tasks = get_daily_checklist()
    for row_idx, t in enumerate(daily_tasks, 2):
        chk = "☑" if t[4] == "Hoàn thành" else "☐"
        type_str = "🔄 Lặp lại hàng ngày" if t[6] == "daily" else "📌 Việc phát sinh"

        ws0.cell(row=row_idx, column=1, value=row_idx-1).alignment = Alignment(horizontal="center")
        ws0.cell(row=row_idx, column=2, value=chk).alignment = Alignment(horizontal="center")
        ws0.cell(row=row_idx, column=3, value=t[1])
        ws0.cell(row=row_idx, column=4, value=type_str).alignment = Alignment(horizontal="center")
        ws0.cell(row=row_idx, column=5, value=t[2] or "Chung").alignment = Alignment(horizontal="center")

        p_cell = ws0.cell(row=row_idx, column=6, value=t[3])
        p_cell.alignment = Alignment(horizontal="center")
        if t[3] in priority_fonts:
            p_cell.font = priority_fonts[t[3]]

        ws0.cell(row=row_idx, column=7, value=t[5] or "").alignment = Alignment(horizontal="center")

        st_cell = ws0.cell(row=row_idx, column=8, value=t[4])
        st_cell.alignment = Alignment(horizontal="center")
        if t[4] in status_fills:
            st_cell.fill = status_fills[t[4]]
            st_cell.font = status_fonts[t[4]]

        ws0.cell(row=row_idx, column=9, value=t[7] or "").alignment = Alignment(horizontal="center")

        for col_num in range(1, 10):
            c = ws0.cell(row=row_idx, column=col_num)
            c.border = thin_border
            if col_num not in [6, 8]:
                c.font = data_font

    # -------------------------------------------------------------
    # TAB 2: DANH SÁCH CÔNG VIỆC TỔNG THỂ
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="📋 Tất cả công việc")
    ws1.views.sheetView[0].showGridLines = True
    
    headers1 = ["Mã ID", "Tên công việc", "Ghi chú / Chi tiết", "Phân loại", "Độ ưu tiên", "Trạng thái", "Hạn chót / Giờ nhắc", "Lặp lại", "Ngày tạo", "Ngày hoàn thành"]
    for col_num, header in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    tasks = get_all_tasks()
    for row_idx, task in enumerate(tasks, 2):
        ws1.cell(row=row_idx, column=1, value=task[0]).alignment = Alignment(horizontal="center")
        ws1.cell(row=row_idx, column=2, value=task[1])
        ws1.cell(row=row_idx, column=3, value=task[2] or "")
        ws1.cell(row=row_idx, column=4, value=task[3] or "Chung").alignment = Alignment(horizontal="center")
        
        p_cell = ws1.cell(row=row_idx, column=5, value=task[4])
        p_cell.alignment = Alignment(horizontal="center")
        if task[4] in priority_fonts:
            p_cell.font = priority_fonts[task[4]]

        st_cell = ws1.cell(row=row_idx, column=6, value=task[5])
        st_cell.alignment = Alignment(horizontal="center")
        if task[5] in status_fills:
            st_cell.fill = status_fills[task[5]]
            st_cell.font = status_fonts[task[5]]

        ws1.cell(row=row_idx, column=7, value=task[6] or "").alignment = Alignment(horizontal="center")
        ws1.cell(row=row_idx, column=8, value="Hàng ngày" if task[9]=='daily' else "Không").alignment = Alignment(horizontal="center")
        ws1.cell(row=row_idx, column=9, value=task[7] or "").alignment = Alignment(horizontal="center")
        ws1.cell(row=row_idx, column=10, value=task[8] or "").alignment = Alignment(horizontal="center")

        for col_num in range(1, 11):
            c = ws1.cell(row=row_idx, column=col_num)
            c.border = thin_border
            if col_num not in [5, 6]:
                c.font = data_font

    # -------------------------------------------------------------
    # TAB 3: NHẬT KÝ ĐÃ LÀM
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="📝 Nhật ký đã làm")
    ws2.views.sheetView[0].showGridLines = True
    
    headers2 = ["STT", "Mã Task Ref", "Phân loại", "Ngày thực hiện", "Thời gian", "Nội dung công việc đã làm"]
    for col_num, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    logs = get_all_work_logs()
    for row_idx, log in enumerate(logs, 2):
        ws2.cell(row=row_idx, column=1, value=row_idx-1).alignment = Alignment(horizontal="center")
        ws2.cell(row=row_idx, column=2, value=log[1] if log[1] else "-").alignment = Alignment(horizontal="center")
        ws2.cell(row=row_idx, column=3, value=log[2] or "Chung").alignment = Alignment(horizontal="center")
        ws2.cell(row=row_idx, column=4, value=log[4]).alignment = Alignment(horizontal="center")
        ws2.cell(row=row_idx, column=5, value=log[5]).alignment = Alignment(horizontal="center")
        ws2.cell(row=row_idx, column=6, value=log[3])

        for col_num in range(1, 7):
            c = ws2.cell(row=row_idx, column=col_num)
            c.border = thin_border
            c.font = data_font

    # -------------------------------------------------------------
    # TAB 4: THỐNG KÊ CHUNG
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="📊 Thống kê chung")
    ws3.views.sheetView[0].showGridLines = True
    
    total_tasks = len(tasks)
    completed_tasks = len([t for t in tasks if t[5] == 'Hoàn thành'])
    pending_tasks = len([t for t in tasks if t[5] in ['Chờ làm', 'Đang nhắc', 'Đang làm']])
    rate = (completed_tasks / total_tasks * 100) if total_tasks > 0 else 0
    
    stat_headers = ["Chỉ số hiệu suất", "Giá trị"]
    for col_num, header in enumerate(stat_headers, 1):
        cell = ws3.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        
    stats_data = [
        ("Tổng số công việc đã tạo", total_tasks),
        ("Số việc đã hoàn thành", completed_tasks),
        ("Số việc đang chờ xử lý", pending_tasks),
        ("Tỷ lệ hoàn thành (%)", f"{rate:.1f}%"),
        ("Tổng số nhật ký công việc đã ghi", len(logs))
    ]
    
    for r_i, (k, v) in enumerate(stats_data, 2):
        c1 = ws3.cell(row=r_i, column=1, value=k)
        c2 = ws3.cell(row=r_i, column=2, value=v)
        c1.font = bold_data_font
        c2.font = bold_data_font
        c2.alignment = Alignment(horizontal="center")
        c1.border = thin_border
        c2.border = thin_border

    # Auto column width
    for sheet in [ws0, ws1, ws2, ws3]:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(EXCEL_PATH)
    return EXCEL_PATH

def sync_all_exports():
    excel_path = export_to_excel()
    try:
        from google_sheets_sync import sync_to_google_sheets
        sync_to_google_sheets()
    except Exception as e:
        print(f"Notice Google Sheets sync: {e}")
    return excel_path

if __name__ == "__main__":
    sync_all_exports()
    print("Exported successfully to Excel file!")
